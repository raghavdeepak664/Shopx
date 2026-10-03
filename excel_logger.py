"""
ShopX Excel & CSV Activity Logger
Automatically records:
1. User Logins and Signups
2. Cart Additions and Updates
3. Purchases and Orders Placed
Saves everything in real-time to both:
- shopx_user_data.xlsx (Multi-sheet formatted Excel workbook)
- shopx_activity_log.csv (Easy import into Google Sheets / Excel)
"""

import os
import csv
import threading
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_FILE = os.path.join(BASE_DIR, "shopx_user_data.xlsx")
CSV_FILE = os.path.join(BASE_DIR, "shopx_activity_log.csv")

_lock = threading.Lock()

# Define headers for each sheet
SHEET_HEADERS = {
    "All Activity Stream": [
        "Timestamp", "Event Type", "Customer Name", "Email / Phone",
        "Item / Details", "Price / Amount (INR)", "IP Address", "Device / Browser"
    ],
    "Logins & Signups": [
        "Timestamp", "Action Type", "Name", "Email / Phone",
        "Status", "IP Address", "Device / Browser"
    ],
    "Cart Activity": [
        "Timestamp", "Customer Name", "Product ID", "Product Title",
        "Brand", "Quantity", "Price Each (INR)", "Total Price (INR)", "IP Address"
    ],
    "Orders & Purchases": [
        "Timestamp", "Order ID", "Customer Name", "Phone",
        "Delivery Address", "City", "Pincode", "Payment Method",
        "Items Count", "Items Summary", "Total Paid (INR)", "Status"
    ]
}

def init_excel_workbook():
    """Create the Excel workbook if it doesn't already exist with styled headers."""
    if os.path.exists(EXCEL_FILE):
        return

    with _lock:
        if os.path.exists(EXCEL_FILE):
            return

        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="131921", end_color="131921", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin', color='D3D3D3'),
            right=Side(style='thin', color='D3D3D3'),
            top=Side(style='thin', color='D3D3D3'),
            bottom=Side(style='thin', color='D3D3D3')
        )

        for sheet_name, headers in SHEET_HEADERS.items():
            ws = wb.create_sheet(title=sheet_name)
            ws.append(headers)

            # Style the header row
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border
            ws.row_dimensions[1].height = 26

            # Default column widths
            for col_idx, col_name in enumerate(headers, 1):
                col_letter = get_column_letter(col_idx)
                ws.column_dimensions[col_letter].width = max(len(col_name) + 6, 18)

        wb.save(EXCEL_FILE)

        # Also initialize the unified CSV file
        if not os.path.exists(CSV_FILE):
            with open(CSV_FILE, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f)
                writer.writerow(SHEET_HEADERS["All Activity Stream"])

def log_event(event_type, customer_name, email_phone, details, amount=0, ip_address="127.0.0.1", user_agent=""):
    """Append event to All Activity Stream sheet and CSV."""
    init_excel_workbook()
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    row_data = [
        timestamp,
        event_type,
        customer_name or "Guest User",
        email_phone or "N/A",
        details or "",
        f"₹{amount:,.2f}" if isinstance(amount, (int, float)) and amount > 0 else (str(amount) if amount else "—"),
        ip_address or "127.0.0.1",
        user_agent or "Web Browser"
    ]

    with _lock:
        # Append to CSV
        try:
            with open(CSV_FILE, "a", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f)
                writer.writerow(row_data)
        except Exception as e:
            print("Error writing to CSV:", e)

        # Append to Excel
        try:
            wb = openpyxl.load_workbook(EXCEL_FILE)
            if "All Activity Stream" in wb.sheetnames:
                ws = wb["All Activity Stream"]
                ws.append(row_data)
                
                # Auto color code by event type
                last_row = ws.max_row
                event_cell = ws.cell(row=last_row, column=2)
                if event_type == "LOGIN":
                    event_cell.fill = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid") # light blue
                elif event_type == "SIGNUP":
                    event_cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # light green
                elif event_type == "ORDER_PLACED":
                    event_cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # light amber
                elif event_type == "ADD_TO_CART":
                    event_cell.fill = PatternFill(start_color="F3E8FF", end_color="F3E8FF", fill_type="solid") # light purple

            wb.save(EXCEL_FILE)
        except Exception as e:
            print("Error writing to Excel:", e)

def log_login_or_signup(action_type, name, email_phone, status="SUCCESS", ip_address="127.0.0.1", user_agent=""):
    """Specific logger for Logins & Signups sheet."""
    init_excel_workbook()
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    row_data = [
        timestamp,
        action_type,
        name or "User",
        email_phone,
        status,
        ip_address,
        user_agent
    ]

    # Also log to main stream
    log_event(
        event_type=action_type,
        customer_name=name,
        email_phone=email_phone,
        details=f"User {action_type.lower()} successfully",
        amount=0,
        ip_address=ip_address,
        user_agent=user_agent
    )

    with _lock:
        try:
            wb = openpyxl.load_workbook(EXCEL_FILE)
            if "Logins & Signups" in wb.sheetnames:
                ws = wb["Logins & Signups"]
                ws.append(row_data)
                last_row = ws.max_row
                action_cell = ws.cell(row=last_row, column=2)
                if action_type == "LOGIN":
                    action_cell.fill = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
                else:
                    action_cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                wb.save(EXCEL_FILE)
        except Exception as e:
            print("Error logging login to Excel:", e)

def log_cart_activity(customer_name, product_id, title, brand, quantity, price, action="ADDED", ip_address="127.0.0.1"):
    """Specific logger for Cart Activity sheet."""
    init_excel_workbook()
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    total_price = (price or 0) * (quantity or 1)

    row_data = [
        timestamp,
        customer_name or "Guest",
        product_id,
        title,
        brand or "N/A",
        quantity,
        price,
        total_price,
        ip_address
    ]

    # Also log to main stream
    log_event(
        event_type="ADD_TO_CART",
        customer_name=customer_name,
        email_phone="",
        details=f"{action}: {title} (x{quantity})",
        amount=total_price,
        ip_address=ip_address
    )

    with _lock:
        try:
            wb = openpyxl.load_workbook(EXCEL_FILE)
            if "Cart Activity" in wb.sheetnames:
                ws = wb["Cart Activity"]
                ws.append(row_data)
                wb.save(EXCEL_FILE)
        except Exception as e:
            print("Error logging cart to Excel:", e)

def log_order_purchase(order_dict, ip_address="127.0.0.1"):
    """Specific logger for Orders & Purchases sheet."""
    init_excel_workbook()
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    items = order_dict.get("items", [])
    items_count = sum(i.get("quantity", 1) for i in items)
    items_summary = "; ".join([f"{i.get('title', 'Item')[:25]} (x{i.get('quantity', 1)})" for i in items])

    row_data = [
        timestamp,
        order_dict.get("order_id", "N/A"),
        order_dict.get("name") or order_dict.get("customer_name") or "Customer",
        order_dict.get("phone") or order_dict.get("customer_phone") or "N/A",
        order_dict.get("address", "N/A"),
        order_dict.get("city", "N/A"),
        order_dict.get("pincode", "N/A"),
        order_dict.get("payment_method", "UPI"),
        items_count,
        items_summary,
        order_dict.get("total", 0),
        order_dict.get("status", "Order Confirmed")
    ]

    # Also log to main stream
    log_event(
        event_type="ORDER_PLACED",
        customer_name=order_dict.get("name") or order_dict.get("customer_name"),
        email_phone=order_dict.get("phone") or order_dict.get("customer_phone"),
        details=f"Order {order_dict.get('order_id')}: {items_summary}",
        amount=order_dict.get("total", 0),
        ip_address=ip_address
    )

    with _lock:
        try:
            wb = openpyxl.load_workbook(EXCEL_FILE)
            if "Orders & Purchases" in wb.sheetnames:
                ws = wb["Orders & Purchases"]
                ws.append(row_data)
                wb.save(EXCEL_FILE)
        except Exception as e:
            print("Error logging order to Excel:", e)

# Test/Verify module initialization on load
init_excel_workbook()
